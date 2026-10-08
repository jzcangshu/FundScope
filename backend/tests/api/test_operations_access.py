from __future__ import annotations

from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.db.base import ImportBatchStatus, SourceType, ValuationStatus
from app.db.models import (
    Fund,
    ImportBatch,
    ImportBatchFile,
    SourceFile,
    UserFundAccess,
    ValuationVersion,
)
from app.imports.access import ImportAccessError
from app.imports.processor import process_import_batch
from tests.imports.test_processor import _valuation_xlsx

from .conftest import seed_pending_version


def _operator(admin, fund_ids=(), *, mode="selected"):
    response = admin.post(
        "/api/v1/users",
        json={
            "username": "scoped-operator",
            "password": "correct horse",
            "role": "operator",
            "fund_access_mode": mode,
            "fund_ids": list(fund_ids),
        },
    )
    assert response.status_code == 201, response.text
    client = TestClient(admin.app)
    login = client.post(
        "/api/v1/auth/login",
        json={
            "username": "scoped-operator",
            "password": "correct horse",
        },
    )
    assert login.status_code == 200
    return client, response.json()["data"]["id"]


def _batch(client):
    response = client.post("/api/v1/imports", json={})
    assert response.status_code == 201
    return response.json()["data"]["id"]


def _upload(client, batch, data, name="valuation.xlsx"):
    return client.post(f"/api/v1/imports/{batch}/files", files={"file": (name, data)})


def _funds(engine):
    with Session(engine) as session:
        first, second = Fund(standard_name="千帆一号"), Fund(standard_name="未授权产品")
        session.add_all([first, second])
        session.commit()
        return first.id, second.id


def test_catalog_writes_and_global_operations_respect_product_access(
    admin_client, app_and_engine
):
    first, second = _funds(app_and_engine[1])
    operator, _ = _operator(admin_client, [first])
    assert (
        operator.patch(
            f"/api/v1/funds/{first}", json={"notes": "authorized"}
        ).status_code
        == 200
    )
    for fund_id, expected in [(first, 200), (second, 404)]:
        assert operator.get(f"/api/v1/funds/{fund_id}/aliases").status_code == expected
        assert (
            operator.get(f"/api/v1/funds/{fund_id}/share-classes").status_code
            == expected
        )
    assert (
        operator.patch(f"/api/v1/funds/{second}", json={"notes": "denied"}).status_code
        == 404
    )
    assert (
        operator.post(
            f"/api/v1/funds/{second}/aliases", json={"alias": "new-alias"}
        ).status_code
        == 404
    )
    assert (
        operator.post(
            f"/api/v1/funds/{second}/share-classes",
            json={"share_code": "A", "share_name": "A"},
        ).status_code
        == 404
    )
    assert (
        operator.post(
            "/api/v1/funds",
            json={
                "standard_name": "new-fund",
                "product_code": "new-code",
                "establishment_date": "2026-01-01",
                "aliases": [{"alias": "new-alias"}],
            },
        ).status_code
        == 403
    )
    for path in [
        "/subjects/mappings",
        "/mail/messages",
        "/mail/settings",
        "/system/operations",
        "/audit-logs",
    ]:
        assert operator.get("/api/v1" + path).status_code == 403, path
    with Session(app_and_engine[1]) as session:
        assert session.get(Fund, second).notes is None


@pytest.mark.parametrize(
    "mode,granted,expected",
    [("selected", True, 1), ("selected", False, 0), ("all", False, 2)],
)
def test_review_queue_and_batch_publication_are_scoped(
    admin_client, app_and_engine, mode, granted, expected
):
    first, first_version = seed_pending_version(app_and_engine[1], fund_name="visible")
    _, second_version = seed_pending_version(app_and_engine[1], fund_name="hidden")
    operator, _ = _operator(admin_client, [first] if granted else [], mode=mode)
    queue = operator.get("/api/v1/reviews").json()
    assert queue["meta"]["total"] == expected
    assert len(queue["data"]) == expected
    if mode == "selected":
        for action, body in [
            ("publish", {"reason": "denied"}),
            ("reject", {"reason": "denied"}),
            ("revoke", {"reason": "denied"}),
            ("restore", {"reason": "denied"}),
        ]:
            assert (
                operator.post(
                    f"/api/v1/valuations/{second_version}/{action}", json=body
                ).status_code
                == 404
            )
        assert (
            operator.post(
                f"/api/v1/reviews/{second_version}/acknowledge",
                json={"allow_publish": True, "note": "denied"},
            ).status_code
            == 404
        )
    with Session(app_and_engine[1]) as session:
        for version_id in [first_version, second_version]:
            session.get(
                ValuationVersion, version_id
            ).status = ValuationStatus.PUBLISHABLE
        session.commit()
    response = operator.post(
        "/api/v1/reviews/batch-publish", json={"reason": "scoped publication"}
    )
    assert response.status_code == 200, response.text
    assert response.json()["data"]["requested"] == expected
    assert response.json()["data"]["published"] == expected
    with Session(app_and_engine[1]) as session:
        if mode == "selected":
            assert (
                session.get(ValuationVersion, second_version).status
                == ValuationStatus.PUBLISHABLE
            )


def test_authorized_upload_is_readable_and_processes(admin_client, app_and_engine):
    first, _ = _funds(app_and_engine[1])
    operator, _ = _operator(admin_client, [first])
    batch = _batch(operator)
    uploaded = _upload(operator, batch, _valuation_xlsx())
    assert uploaded.status_code == 201, uploaded.text
    source_id = uploaded.json()["data"]["id"]
    assert operator.get(f"/api/v1/imports/{batch}").status_code == 200
    assert (
        operator.get(f"/api/v1/imports/{batch}/source/{source_id}").status_code == 200
    )
    assert operator.post(f"/api/v1/imports/{batch}/complete").status_code == 200
    with Session(app_and_engine[1]) as session:
        result = process_import_batch(session, batch, app_and_engine[0].state.settings)
        assert result.processed_files == 1
        assert {v.fund_id for v in session.scalars(select(ValuationVersion))} == {first}


def test_denied_duplicate_and_unknown_workbooks_do_not_link_or_reveal_files(
    admin_client, app_and_engine
):
    _, second = _funds(app_and_engine[1])
    existing_batch = _batch(admin_client)
    uploaded = _upload(
        admin_client, existing_batch, _valuation_xlsx(), "private-name.xlsx"
    )
    assert uploaded.status_code == 201
    operator, _ = _operator(admin_client, [second])
    batch = _batch(operator)
    denied = _upload(operator, batch, _valuation_xlsx(), "pretend-authorized.xlsx")
    assert denied.status_code == 403
    assert "private-name" not in denied.text
    assert "file_hash" not in denied.text
    assert operator.get(f"/api/v1/imports/{existing_batch}").status_code == 404
    assert (
        operator.get(f"/api/v1/imports/{existing_batch}/validations").status_code == 404
    )
    workbook = Workbook()
    workbook.active.append(["unknown product"])
    data = BytesIO()
    workbook.save(data)
    assert _upload(operator, batch, data.getvalue()).status_code == 403
    with Session(app_and_engine[1]) as session:
        assert session.scalar(select(func.count(SourceFile.id))) == 1
        assert session.scalar(select(func.count(ImportBatchFile.id))) == 1
        assert session.get(ImportBatch, batch).file_count == 0


def test_unknown_and_mixed_batches_are_hidden_including_null_source_identity(
    admin_client, app_and_engine
):
    first, second = _funds(app_and_engine[1])
    operator, _ = _operator(admin_client, [first])
    empty = _batch(operator)
    with Session(app_and_engine[1]) as session:
        batches = []
        for index, fund_ids in enumerate([[first], [second], [None], [first, second]]):
            batch = ImportBatch(source_type=SourceType.UPLOAD, file_count=len(fund_ids))
            session.add(batch)
            session.flush()
            for slot, fund_id in enumerate(fund_ids):
                source = SourceFile(
                    original_filename=f"{index}-{slot}.xlsx",
                    file_hash=f"{index * 10 + slot:064d}",
                    file_size=1,
                    file_extension=".xlsx",
                    source_type=SourceType.UPLOAD,
                    object_name=f"{index}-{slot}.xlsx",
                    recognized_fund_id=fund_id,
                )
                session.add(source)
                session.flush()
                session.add(
                    ImportBatchFile(batch_id=batch.id, source_file_id=source.id)
                )
            batches.append(batch.id)
        session.commit()
    response = operator.get("/api/v1/imports")
    assert response.status_code == 200, response.text
    assert {batch["id"] for batch in response.json()["data"]} == {empty, batches[0]}
    assert response.json()["meta"]["total"] == 2
    for batch_id in batches[1:]:
        assert operator.get(f"/api/v1/imports/{batch_id}").status_code == 404


def test_permissions_are_rechecked_after_upload_before_background_writes(
    admin_client, app_and_engine
):
    first, _ = _funds(app_and_engine[1])
    operator, user_id = _operator(admin_client, [first])
    batch = _batch(operator)
    assert _upload(operator, batch, _valuation_xlsx()).status_code == 201
    assert operator.post(f"/api/v1/imports/{batch}/complete").status_code == 200
    with Session(app_and_engine[1]) as session:
        session.delete(session.get(UserFundAccess, (user_id, first)))
        session.commit()
    assert operator.get(f"/api/v1/imports/{batch}").status_code == 404
    with Session(app_and_engine[1]) as session:
        with pytest.raises(ImportAccessError):
            process_import_batch(session, batch, app_and_engine[0].state.settings)
        assert session.scalar(select(func.count(ValuationVersion.id))) == 0
        session.rollback()
        assert session.get(ImportBatch, batch).status == ImportBatchStatus.QUEUED


def test_extra_sheet_cannot_hide_an_unauthorized_product(admin_client, app_and_engine):
    first, _ = _funds(app_and_engine[1])
    operator, _ = _operator(admin_client, [first])
    workbook = load_workbook(BytesIO(_valuation_xlsx()))
    workbook.create_sheet("hidden data").append(["未授权产品", 123456])
    output = BytesIO()
    workbook.save(output)
    assert _upload(operator, _batch(operator), output.getvalue()).status_code == 403


def test_admin_import_does_not_certify_mixed_original_files_for_scoped_downloads(
    admin_client, app_and_engine
):
    first, _ = _funds(app_and_engine[1])
    operator, _ = _operator(admin_client, [first])
    workbook = load_workbook(BytesIO(_valuation_xlsx()))
    workbook.create_sheet("other product").append(["未授权产品", 123456])
    output = BytesIO()
    workbook.save(output)
    batch = _batch(admin_client)
    uploaded = _upload(admin_client, batch, output.getvalue())
    assert uploaded.status_code == 201
    source_id = uploaded.json()["data"]["id"]
    assert admin_client.post(f"/api/v1/imports/{batch}/complete").status_code == 200
    with Session(app_and_engine[1]) as session:
        result = process_import_batch(session, batch, app_and_engine[0].state.settings)
        assert result.created_versions
        assert session.get(SourceFile, source_id).recognized_fund_id is None
        version_status = session.get(
            ValuationVersion, result.created_versions[0]
        ).status
        session.commit()
    assert operator.get(f"/api/v1/imports/{batch}").status_code == 404
    assert (
        operator.get(f"/api/v1/imports/{batch}/source/{source_id}").status_code == 404
    )
    queue = operator.get("/api/v1/reviews", params={"status": version_status}).json()
    assert queue["meta"]["total"] == 1
    assert queue["data"][0]["source_filename"] is None
    assert queue["data"][0]["import_batch_id"] is None


def test_nav_history_imports_enforce_contents_and_reject_mixed_products(
    admin_client, app_and_engine
):
    first, _ = _funds(app_and_engine[1])
    operator, _ = _operator(admin_client, [first])

    def nav(names):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(
            [
                "产品名称",
                "产品编码",
                "净值日期",
                "单位净值",
                "累计净值",
                "资产净值",
                "资产总值",
                "资产份额",
            ]
        )
        for name in names:
            sheet.append([name, "", "2026-08-25", 1, 1, 100, 100, 100])
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    batch = _batch(operator)
    assert _upload(operator, batch, nav(["千帆一号"])).status_code == 201
    assert operator.post(f"/api/v1/imports/{batch}/complete").status_code == 200
    with Session(app_and_engine[1]) as session:
        result = process_import_batch(session, batch, app_and_engine[0].state.settings)
        assert result.created_versions
        assert {v.fund_id for v in session.scalars(select(ValuationVersion))} == {first}
        session.commit()
    assert (
        _upload(operator, _batch(operator), nav(["千帆一号", "未授权产品"])).status_code
        == 403
    )
