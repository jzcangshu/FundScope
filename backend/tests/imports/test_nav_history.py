"""NAV history parser and import chain tests.

The provider layouts below mirror the real 产品净值表 files: each helper
builds a minimal workbook of one layout, and the processor test runs the
full import chain (parse → persist → validate → publish) against them.
"""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.auth.service import AuthService
from app.db.base import SourceType, ValuationStatus
from app.db.models import (
    Fund,
    FundDailySnapshot,
    ShareClass,
    ShareClassDailySnapshot,
    ValuationVersion,
)
from app.imports.processor import process_import_batch
from app.imports.service import ImportService
from app.parser.nav_history import (
    NavHistoryEntry,
    NavHistoryParseError,
    parse_nav_history,
)


def _xlsx(header: list[object], rows: list[list[object]]) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _custodian_list() -> BytesIO:
    # 云川一号/云川二号/明曜三号 style, with optional share-class rows.
    return _xlsx(
        [
            "日期",
            "基金名称",
            "基金代码",
            "资产净值(元)",
            "资产份额(份)",
            "资产总值(元)",
            "单位净值",
            "累计单位净值",
            "净值波动率(%)",
        ],
        [
            [
                "2026-01-05",
                "测试云川一号私募证券投资基金",
                "FS0001",
                "10,200,333.90",
                10200000,
                "10,203,333.90",
                "1.0000",
                "1.0000",
                0.0,
            ],
            [
                "2026-01-06",
                "测试云川一号私募证券投资基金",
                "FS0001",
                "10,199,742.23",
                10200000,
                "10,203,901.56",
                "0.9999",
                "0.9999",
                -0.01,
            ],
        ],
    )


def _ta_series() -> BytesIO:
    # 启航二号 style: 产品名称 产品编码 净值日期 单位净值 累计净值 资产净值 资产总值 资产份额
    return _xlsx(
        [
            "产品名称",
            "产品编码",
            "净值日期",
            "单位净值",
            "累计净值",
            "资产净值",
            "资产总值",
            "资产份额",
        ],
        [
            [
                "测试启航二号私募证券投资基金",
                "FS0008",
                "2026-01-05",
                "1.0001",
                "1.0001",
                "41,003,147.73",
                "41,004,992.73",
                "41,000,000.00",
            ],
            [
                "测试启航二号私募证券投资基金",
                "FS0008",
                "2026-01-06",
                "1.0001",
                "1.0001",
                "41,003,588.52",
                "41,006,837.73",
                "41,000,000.00",
            ],
        ],
    )


def _broker_download() -> BytesIO:
    # 紫微/磐石二号 style: no product name column at all.
    return _xlsx(
        [
            "产品代码",
            "净值日增长率",
            "上一年度日均资产净值(元)",
            "本年度日均资产净值(元)",
            "资产净值预警天数",
            "净值日期",
            "单位净值(元/份)",
            "累计单位净值(元/份)",
            "资产净值(元)",
        ],
        [
            [
                "FS0005",
                "0.0000%",
                "",
                "",
                "0",
                "2026-01-05",
                "1.000",
                "1.000",
                "31,402,693.00",
            ],
            [
                "FS0005",
                "0.2000%",
                "",
                "",
                "0",
                "2026-01-08",
                "1.002",
                "1.002",
                "31,460,450.99",
            ],
        ],
    )


def _internal_table() -> BytesIO:
    # 星辰 style: full balance sheet.
    return _xlsx(
        [
            "产品名称",
            "产品代码",
            "净值日期",
            "单位净值",
            "累计净值",
            "每日净值增长率",
            "累计净值增长率",
            "基金资产份额",
            "基金资产净值",
            "总资产合计",
            "总负债合计",
        ],
        [
            [
                "测试星辰私募证券投资基金",
                "FS0009",
                "2026-01-05",
                "1.0001",
                "1.0001",
                "0.0100%",
                "0.0100%",
                "2,000,000.00",
                "2,000,136.10",
                "2,000,136.10",
                "0.00",
            ],
            [
                "测试星辰私募证券投资基金",
                "FS0009",
                "2026-01-06",
                "1.0000",
                "1.0000",
                "-0.0100%",
                "0.0000%",
                "2,000,000.00",
                "2,000,064.86",
                "2,000,136.10",
                "71.24",
            ],
        ],
    )


def _ta_export() -> BytesIO:
    # 慧达二号 style: 日期 产品代码 协会备案代码 产品名称 资产净值 基金总净值 资产份额 单位净值 累计单位净值
    return _xlsx(
        [
            "日期",
            "产品代码",
            "协会备案代码",
            "产品名称",
            "资产净值",
            "基金总净值",
            "资产份额",
            "单位净值",
            "累计单位净值",
        ],
        [
            [
                "2026/01/05",
                "XY8410",
                "FS0011",
                "测试慧达二号私募证券投资基金",
                26000938.89,
                26000938.89,
                26000000,
                1.0,
                1.0,
            ],
            [
                "2026/01/06",
                "XY8410",
                "FS0011",
                "测试慧达二号私募证券投资基金",
                26001144.05,
                26001144.05,
                26000000,
                1.0001,
                1.0001,
            ],
        ],
    )


def _custodian_recon() -> BytesIO:
    # 明曜一号 style: share classes carry a 类 suffix and a separate
    # 分级基金代码; extra custodian metadata columns are ignored.
    return _xlsx(
        [
            "基金代码",
            "产品名称",
            "业务日期",
            "分级基金代码",
            "单位净值",
            "累计单位净值",
            "托管人",
            "成立日期",
            "清盘日期",
            "状态",
            "资产净值",
        ],
        [
            [
                "FS0012",
                "测试明曜私募证券投资基金B类",
                "2026-01-05",
                "XR911B",
                "0.0000",
                "0.0000",
                "示例证券股份有限公司",
                "2022-11-17",
                "",
                "核对一致",
                "0",
            ],
            [
                "FS0012",
                "测试明曜私募证券投资基金A类",
                "2026-01-05",
                "XR911A",
                "1.0001",
                "1.0001",
                "示例证券股份有限公司",
                "2022-11-17",
                "",
                "核对一致",
                "7,300,379.19",
            ],
            [
                "FS0012",
                "测试明曜私募证券投资基金",
                "2026-01-05",
                "FS0012",
                "1.0001",
                "1.0001",
                "示例证券股份有限公司",
                "2022-11-17",
                "",
                "核对一致",
                "7,300,379.19",
            ],
        ],
    )


def _ta_confirmed() -> BytesIO:
    # 慧达一号 style: three rows per date (fund / A / B).
    return _xlsx(
        [
            "确认状态",
            "产品名称",
            "协会备案编码",
            "TA代码",
            "净值日期",
            "单位净值",
            "累计单位净值",
            "资产净值",
            "资产份额",
            "总资产",
            "总负债",
        ],
        [
            [
                "净值已确认",
                "测试慧达一号私募证券投资基金",
                "FS0010",
                "ZB1568",
                "2026-01-05",
                "1.0000",
                "1.0000",
                "12,300,068.31",
                "12,300,034.14",
                "12,300,068.31",
                "0.00",
            ],
            [
                "净值已确认",
                "测试慧达一号私募证券投资基金A",
                "FS0010",
                "ZB1568",
                "2026-01-05",
                "1.0000",
                "1.0000",
                "9,000,049.99",
                "9,000,024.99",
                "12,300,068.31",
                "0.00",
            ],
            [
                "净值已确认",
                "测试慧达一号私募证券投资基金B",
                "FS0010",
                "ZB1569",
                "2026-01-05",
                "1.0000",
                "1.0000",
                "3,300,018.32",
                "3,300,009.15",
                "12,300,068.31",
                "0.00",
            ],
        ],
    )


def _browse_table() -> BytesIO:
    # 云川三号 style: title rows above the header, 中文日期.
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["基金资产净值浏览表"])
    sheet.append(["长江证券资产托管部_基金资产净值浏览表_专用表"])
    sheet.append(["日期：2026-01-04 至 2026-01-06"])
    sheet.append(
        [
            "账套名称",
            "产品代码",
            "日期",
            "资产净值",
            "资产总值",
            "资产份额",
            "单位净值",
            "累计净值",
            "净值是否确认",
        ]
    )
    sheet.append(
        [
            "测试云川三号私募证券投资基金",
            "FS0003",
            "2026年01月06日",
            "23,000,330.29",
            "23,001,111.68",
            "23,000,000.00",
            "1.0000",
            "1.0000",
            "已确认",
        ]
    )
    sheet.append(
        [
            "测试云川三号私募证券投资基金",
            "FS0003",
            "2026年01月05日",
            "23,000,651.68",
            "23,000,651.68",
            "23,000,000.00",
            "1.0000",
            "1.0000",
            "已确认",
        ]
    )
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _write_tmp(tmp_path: Path, name: str, payload: BytesIO) -> Path:
    path = tmp_path / name
    path.write_bytes(payload.getvalue())
    return path


def test_parse_custodian_list_layout(tmp_path: Path) -> None:
    entries = parse_nav_history(_write_tmp(tmp_path, "a.xlsx", _custodian_list()))
    assert len(entries) == 2
    assert entries[0].product_name == "测试云川一号私募证券投资基金"
    assert entries[0].product_code == "FS0001"
    assert entries[0].valuation_date == date(2026, 1, 5)
    assert entries[0].unit_nav == Decimal("1.0000")
    assert entries[0].net_asset_value == Decimal("10200333.90")
    assert entries[0].total_assets == Decimal("10203333.90")
    assert entries[0].shares_outstanding == Decimal(10200000)
    assert entries[0].daily_return == Decimal(0)
    assert entries[1].daily_return == Decimal("-0.0001")


def test_parse_ta_series_layout(tmp_path: Path) -> None:
    entries = parse_nav_history(_write_tmp(tmp_path, "a.xlsx", _ta_series()))
    assert entries[0].product_name == "测试启航二号私募证券投资基金"
    assert entries[0].product_code == "FS0008"
    assert entries[0].total_assets == Decimal("41004992.73")
    assert entries[0].daily_return is None


def test_parse_broker_download_layout(tmp_path: Path) -> None:
    entries = parse_nav_history(_write_tmp(tmp_path, "a.xlsx", _broker_download()))
    assert entries[0].product_name == ""
    assert entries[0].product_code == "FS0005"
    assert entries[0].unit_nav == Decimal("1.000")
    assert entries[1].daily_return == Decimal("0.002")
    assert entries[0].total_assets is None


def test_parse_internal_table_layout(tmp_path: Path) -> None:
    entries = parse_nav_history(_write_tmp(tmp_path, "a.xlsx", _internal_table()))
    assert entries[0].total_liabilities == Decimal("0.00")
    assert entries[1].total_liabilities == Decimal("71.24")
    assert entries[1].daily_return == Decimal("-0.0001")


def test_parse_ta_export_layout(tmp_path: Path) -> None:
    entries = parse_nav_history(_write_tmp(tmp_path, "a.xlsx", _ta_export()))
    assert entries[0].valuation_date == date(2026, 1, 5)
    assert entries[0].product_code == "XY8410"
    assert entries[0].net_asset_value == Decimal("26000938.89")


def test_parse_ta_confirmed_layout_aggregates_share_classes(
    tmp_path: Path,
) -> None:
    entries = parse_nav_history(_write_tmp(tmp_path, "a.xlsx", _ta_confirmed()))
    assert len(entries) == 1
    entry = entries[0]
    assert entry.product_name == "测试慧达一号私募证券投资基金"
    assert entry.net_asset_value == Decimal("12300068.31")
    assert {item.share_code for item in entry.share_classes} == {"ZB1568", "ZB1569"}
    class_assets = sum((item.net_assets for item in entry.share_classes), Decimal(0))
    assert class_assets == entry.net_asset_value


def test_parse_browse_table_layout_skips_title_rows(tmp_path: Path) -> None:
    entries = parse_nav_history(_write_tmp(tmp_path, "a.xlsx", _browse_table()))
    assert [item.valuation_date for item in entries] == [
        date(2026, 1, 5),
        date(2026, 1, 6),
    ]
    assert entries[0].product_name == "测试云川三号私募证券投资基金"
    assert entries[0].product_code == "FS0003"


def test_parse_custodian_recon_layout_with_class_suffix(tmp_path: Path) -> None:
    entries = parse_nav_history(_write_tmp(tmp_path, "a.xlsx", _custodian_recon()))
    assert len(entries) == 1
    entry = entries[0]
    assert entry.product_name == "测试明曜私募证券投资基金"
    assert entry.product_code == "FS0012"
    assert entry.unit_nav == Decimal("1.0001")
    assert entry.net_asset_value == Decimal("7300379.19")
    assert {item.share_code for item in entry.share_classes} == {"XR911A", "XR911B"}


def _custodian_merged() -> BytesIO:
    # 云川一号/云川二号/明曜三号 merged multi-day file: same list, no
    # 净值波动率 column.
    return _xlsx(
        [
            "日期",
            "基金名称",
            "基金代码",
            "资产净值（元）",
            "资产份额（份）",
            "单位净值",
            "累计单位净值",
            "资产总值（元）",
        ],
        [
            [
                "2026-09-14",
                "上海测试云川一号私募证券投资基金",
                "FS0001",
                "45446784.83",
                "48816452.06",
                "0.9310",
                "1.7555",
                "45618664.79",
            ],
            [
                "2026-09-15",
                "上海测试云川一号私募证券投资基金",
                "FS0001",
                "45556917.77",
                "48816452.06",
                "0.9332",
                "1.7577",
                "45733932.09",
            ],
        ],
    )


def _fund_daily_table() -> BytesIO:
    # 启航二号 基金每日净值表: 资产代码/资产名称 and 资产份额净值 labels.
    return _xlsx(
        [
            "日期",
            "资产代码",
            "资产名称",
            "资产份额净值(元)",
            "资产份额累计净值(元)",
            "总份额",
            "资产总值(元)",
        ],
        [
            [
                "2026-09-16",
                "FS0008",
                "测试启航二号私募证券投资基金",
                "1.3527",
                "1.3527",
                "37,156,587.85",
                "50,447,417.51",
            ],
        ],
    )


def _broker_announcement() -> BytesIO:
    # 青岩稳健/紫微/磐石二号 daily announcement: 产品总份额 / 产品资产净值 /
    # 产品总资产 labels, inline strings with a bogus dimension in production.
    return _xlsx(
        [
            "净值日期",
            "产品代码",
            "产品名称",
            "单位净值",
            "累计单位净值",
            "产品总份额",
            "产品资产净值",
            "产品总资产",
        ],
        [
            [
                "2026-09-15",
                "FS0007",
                "测试磐石二号私募证券投资基金",
                "2.841",
                "2.841",
                "9,407,609.4500",
                "26,722,986.27",
                "26,927,218.16",
            ],
        ],
    )


def _ta_confirmed_export() -> BytesIO:
    # 慧达一号 产品净值 export: bilingual header annotations in parentheses.
    return _xlsx(
        [
            "日期\n（NAV As Of Date）",
            "产品名称\n（Fund Name）",
            "单位净值\n（NAV/Share）",
            "累计单位净值\n（Accumulated NAV/Share）",
            "资产份额\n（Fund Shares）",
            "资产净值\n（Net Asset Value）",
            "总资产\n（Total Assets Value）",
            "产品代码\n（TA Code）",
            "协会备案编码\n（Fund Filling Code）",
        ],
        [
            [
                "2026-09-16",
                "测试慧达一号私募证券投资基金",
                "1.2901",
                "1.7913",
                "30,626,495.50",
                "39,510,608.21",
                "39,730,932.88",
                "ZB1568",
                "FS0010",
            ],
        ],
    )


def _outsourced_review() -> BytesIO:
    # 慧达二号 补发净值表: 8-digit dates, 备案代码 and 产品代码 both present.
    return _xlsx(
        [
            "估值日期",
            "协会备案代码",
            "产品代码",
            "产品全称",
            "单位净值",
            "累计单位净值",
            "资产净值",
            "资产份额",
            "净值状态",
        ],
        [
            [
                "20260914",
                "FS0011",
                "XY8410",
                "测试慧达二号私募证券投资基金",
                "0.893",
                "1.2975",
                "49750514.24",
                "55708728.15",
                "外包确认,托管确认",
            ],
            [
                "20260915",
                "FS0011",
                "XY8410",
                "测试慧达二号私募证券投资基金",
                "0.8865",
                "1.2909",
                "49383480.56",
                "55708728.15",
                "外包确认,托管确认",
            ],
        ],
    )


def _outsourced_review_compact() -> BytesIO:
    # 天衡上将 补发净值表: same export without 资产净值/资产份额 columns.
    return _xlsx(
        [
            "估值日期",
            "协会备案代码",
            "产品代码",
            "产品全称",
            "单位净值",
            "累计单位净值",
            "净值状态",
        ],
        [
            [
                "20260914",
                "SATP94",
                "XY8306",
                "测试天衡上将私募证券投资基金",
                "2.8363",
                "2.8363",
                "外包确认,托管确认",
            ],
        ],
    )


def _custodian_announcement() -> BytesIO:
    # 云川三号 daily announcement: two-row header (账套名称/基金代码 over the
    # NAV block) with a stray 资产净值 label under 基金代码, date in the title.
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["资产净值公告"])
    sheet.append(["长江证券__测试云川三号私募证券投资基金__专用表"])
    sheet.append(["日期：2026-09-16"])
    sheet.append(
        [
            "经定向资产管理计划托管人核准，截至2026-09-16,以下定向资产管理计划资产净值如下："
        ]
    )
    sheet.append(["账套名称", "基金代码", "净值情况", "", "", "", ""])
    sheet.append(
        ["", "资产净值", "单位净值", "累计单位净值", "资产份额", "资产净值", "资产总值"]
    )
    sheet.append(
        [
            "测试云川三号私募证券投资基金",
            "FS0003",
            "0.8856",
            "0.8856",
            "29778468.49",
            "26370446.66",
            "26438257.54",
        ]
    )
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _key_value_announcement() -> BytesIO:
    # 星辰 daily 净值信息: label rows only, date in the title block.
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["资产净值公告"])
    sheet.append(["测试星辰私募证券投资基金专用表"])
    sheet.append(["2026年09月16日"])
    sheet.append(["经基金托管人核准，截至2026年09月16日,以下基金资产净值如下："])
    sheet.append(["单位：人民币元"])
    sheet.append(["基金代码：", "FS0009"])
    sheet.append(["基金名称：", "测试星辰私募证券投资基金"])
    sheet.append(["基金资产净值：", "8,464,943.23"])
    sheet.append(["基金资产份额：", "8,449,467.78"])
    sheet.append(["基金份额净值：", "1.0018"])
    sheet.append(["基金份额累计净值：", "1.0259"])
    sheet.append(["总层面资产合计：", "8,489,619.02"])
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _key_value_share_classes() -> BytesIO:
    # 明曜 示例道富 announcement: label rows indented behind a blank column,
    # one value column per share class.
    rows: list[list[object]] = [
        ["资产净值公告"],
        ["示例道富____专用表"],
        ["2026-09-14"],
        [
            "经定向资产管理计划托管人核准，截至2026-09-14,以下定向资产管理计划资产净值如下："
        ],
        ["单位：人民币元"],
        [" 基金代码：", "", "FS0012", "XR911A", "XR911B"],
        [
            " 基金名称：",
            "",
            "测试明曜私募证券投资基金",
            "测试明曜私募证券投资基金_A类",
            "测试明曜私募证券投资基金_B类",
        ],
        [" 基金资产净值：", "", "50154543.89", "42855615.78", "7298928.11"],
        [" 基金资产份额：", "", "13908146.7", "11884490.82", "2023655.88"],
        [" 基金份额净值：", "", "3.6061", "3.606", "3.6068"],
        [" 基金份额累计净值：", "", "3.6061", "3.606", "3.6068"],
        [
            " 管理人：",
            "",
            "示例资产管理中心（有限合伙）",
            "示例资产管理中心（有限合伙）",
            "示例资产管理中心（有限合伙）",
        ],
    ]
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def test_parse_custodian_merged_layout(tmp_path: Path) -> None:
    entries = parse_nav_history(_write_tmp(tmp_path, "a.xlsx", _custodian_merged()))
    assert [item.valuation_date for item in entries] == [
        date(2026, 9, 14),
        date(2026, 9, 15),
    ]
    assert entries[1].product_code == "FS0001"
    assert entries[1].net_asset_value == Decimal("45556917.77")
    assert entries[1].total_assets == Decimal("45733932.09")
    assert entries[1].shares_outstanding == Decimal("48816452.06")


def test_parse_fund_daily_table_layout(tmp_path: Path) -> None:
    entries = parse_nav_history(_write_tmp(tmp_path, "a.xlsx", _fund_daily_table()))
    assert entries[0].product_code == "FS0008"
    assert entries[0].product_name == "测试启航二号私募证券投资基金"
    assert entries[0].unit_nav == Decimal("1.3527")
    assert entries[0].cumulative_unit_nav == Decimal("1.3527")
    assert entries[0].total_assets == Decimal("50447417.51")
    assert entries[0].net_asset_value is None


def test_parse_broker_announcement_layout(tmp_path: Path) -> None:
    entries = parse_nav_history(_write_tmp(tmp_path, "a.xlsx", _broker_announcement()))
    assert entries[0].product_code == "FS0007"
    assert entries[0].valuation_date == date(2026, 9, 15)
    assert entries[0].unit_nav == Decimal("2.841")
    assert entries[0].net_asset_value == Decimal("26722986.27")
    assert entries[0].total_assets == Decimal("26927218.16")
    assert entries[0].shares_outstanding == Decimal("9407609.4500")


def test_parse_ta_confirmed_export_with_bilingual_header(tmp_path: Path) -> None:
    entries = parse_nav_history(_write_tmp(tmp_path, "a.xlsx", _ta_confirmed_export()))
    assert len(entries) == 1
    entry = entries[0]
    assert entry.product_name == "测试慧达一号私募证券投资基金"
    assert entry.product_code == "FS0010"
    assert entry.valuation_date == date(2026, 9, 16)
    assert entry.unit_nav == Decimal("1.2901")
    assert entry.total_assets == Decimal("39730932.88")


def test_parse_outsourced_review_layout(tmp_path: Path) -> None:
    entries = parse_nav_history(_write_tmp(tmp_path, "a.xlsx", _outsourced_review()))
    assert [item.valuation_date for item in entries] == [
        date(2026, 9, 14),
        date(2026, 9, 15),
    ]
    assert entries[0].product_name == "测试慧达二号私募证券投资基金"
    # The filing code keys the fund catalogue, so it wins over the TA code.
    assert entries[0].product_code == "FS0011"
    assert entries[0].net_asset_value == Decimal("49750514.24")
    assert entries[0].shares_outstanding == Decimal("55708728.15")


def test_parse_outsourced_review_compact_layout(tmp_path: Path) -> None:
    entries = parse_nav_history(
        _write_tmp(tmp_path, "a.xlsx", _outsourced_review_compact())
    )
    assert entries[0].product_code == "SATP94"
    assert entries[0].valuation_date == date(2026, 9, 14)
    assert entries[0].unit_nav == Decimal("2.8363")
    assert entries[0].net_asset_value is None


def test_parse_custodian_announcement_two_row_header(tmp_path: Path) -> None:
    entries = parse_nav_history(
        _write_tmp(tmp_path, "a.xlsx", _custodian_announcement())
    )
    assert len(entries) == 1
    entry = entries[0]
    assert entry.product_name == "测试云川三号私募证券投资基金"
    assert entry.product_code == "FS0003"
    assert entry.valuation_date == date(2026, 9, 16)
    assert entry.unit_nav == Decimal("0.8856")
    assert entry.cumulative_unit_nav == Decimal("0.8856")
    assert entry.net_asset_value == Decimal("26370446.66")
    assert entry.total_assets == Decimal("26438257.54")
    assert entry.shares_outstanding == Decimal("29778468.49")


def test_parse_key_value_announcement_layout(tmp_path: Path) -> None:
    entries = parse_nav_history(
        _write_tmp(tmp_path, "a.xlsx", _key_value_announcement())
    )
    assert len(entries) == 1
    entry = entries[0]
    assert entry.product_name == "测试星辰私募证券投资基金"
    assert entry.product_code == "FS0009"
    assert entry.valuation_date == date(2026, 9, 16)
    assert entry.unit_nav == Decimal("1.0018")
    assert entry.cumulative_unit_nav == Decimal("1.0259")
    assert entry.net_asset_value == Decimal("8464943.23")
    assert entry.total_assets == Decimal("8489619.02")
    assert entry.shares_outstanding == Decimal("8449467.78")
    assert entry.share_classes == ()


def test_parse_key_value_announcement_share_classes(tmp_path: Path) -> None:
    entries = parse_nav_history(
        _write_tmp(tmp_path, "a.xlsx", _key_value_share_classes())
    )
    assert len(entries) == 1
    entry = entries[0]
    assert entry.product_name == "测试明曜私募证券投资基金"
    assert entry.product_code == "FS0012"
    assert entry.valuation_date == date(2026, 9, 14)
    assert entry.unit_nav == Decimal("3.6061")
    assert entry.net_asset_value == Decimal("50154543.89")
    assert entry.shares_outstanding == Decimal("13908146.7")
    assert {item.share_code for item in entry.share_classes} == {
        "XR911A",
        "XR911B",
    }
    class_assets = sum((item.net_assets for item in entry.share_classes), Decimal(0))
    assert class_assets == entry.net_asset_value


def test_parse_nav_history_rejects_unknown_workbook(tmp_path: Path) -> None:
    try:
        parse_nav_history(
            _write_tmp(tmp_path, "a.xlsx", _xlsx(["随便", "表头"], [[1, 2]]))
        )
    except NavHistoryParseError:
        pass
    else:
        raise AssertionError("expected NavHistoryParseError")


def _nav_entry(**overrides: object) -> NavHistoryEntry:
    defaults: dict[str, object] = {
        "product_name": "测试慧达二号私募证券投资基金",
        "valuation_date": date(2026, 1, 5),
        "product_code": "XY8410",
        "unit_nav": Decimal("1.0000"),
        "cumulative_unit_nav": Decimal("1.0000"),
        "net_asset_value": Decimal(1000000),
        "total_assets": Decimal(1000000),
    }
    return NavHistoryEntry(**{**defaults, **overrides})  # type: ignore[arg-type]


def test_processor_imports_nav_history_and_publishes(
    app_and_engine: tuple[object, object],
    tmp_path: Path,
) -> None:
    app, engine = app_and_engine
    with Session(engine) as session:
        actor = AuthService(session).initialize_admin("admin", "correct horse").user
        fund = Fund(standard_name="测试慧达二号私募证券投资基金", product_code="XY8410")
        session.add(fund)
        session.flush()
        service = ImportService.from_settings(session, app.state.settings)
        batch = service.create_batch(SourceType.UPLOAD, actor.id)
        service.receive_upload(
            batch.id,
            "测试慧达二号净值表.xlsx",
            _ta_export(),
            actor.id,
        )
        service.complete_batch(batch.id, actor.id)
        session.commit()

        result = process_import_batch(session, batch.id, app.state.settings)
        session.commit()

        assert result.processed_files == 1
        assert result.published_files == 1
        assert len(result.created_versions) == 2
        versions = list(
            session.scalars(
                select(ValuationVersion)
                .where(ValuationVersion.fund_id == fund.id)
                .order_by(ValuationVersion.valuation_date)
            )
        )
        assert all(version.status == ValuationStatus.PUBLISHED for version in versions)
        first = session.scalar(
            select(FundDailySnapshot).where(
                FundDailySnapshot.valuation_version_id == versions[0].id
            )
        )
        assert first is not None
        assert first.unit_nav == Decimal("1.0")
        # The xlsx cell is a native float; SQLite round-trips it with binary
        # noise, so compare within a cent instead of exactly.
        assert abs(first.net_asset_value - Decimal("26000938.89")) < Decimal("0.01")
        # 慧达二号 style has no daily return; the previous-day NAV is filled
        # in by the processor for downstream return computation.
        second = session.scalar(
            select(FundDailySnapshot).where(
                FundDailySnapshot.valuation_version_id == versions[1].id
            )
        )
        assert second is not None
        assert second.previous_unit_nav == Decimal("1.0")


def test_processor_resolves_nav_fund_by_code_when_name_absent(
    app_and_engine: tuple[object, object],
    tmp_path: Path,
) -> None:
    app, engine = app_and_engine
    with Session(engine) as session:
        actor = AuthService(session).initialize_admin("admin", "correct horse").user
        fund = Fund(standard_name="测试紫微私募证券投资基金", product_code="FS0005")
        session.add(fund)
        session.flush()
        service = ImportService.from_settings(session, app.state.settings)
        batch = service.create_batch(SourceType.UPLOAD, actor.id)
        service.receive_upload(
            batch.id, "FS0005_净值数据.xlsx", _broker_download(), actor.id
        )
        service.complete_batch(batch.id, actor.id)
        session.commit()

        result = process_import_batch(session, batch.id, app.state.settings)
        session.commit()

        assert result.processed_files == 1
        versions = list(
            session.scalars(
                select(ValuationVersion).where(ValuationVersion.fund_id == fund.id)
            )
        )
        assert len(versions) == 2
        # No balance sheet columns: the check is skipped, not failed.
        assert all(version.status == ValuationStatus.PUBLISHED for version in versions)


def test_processor_imports_share_class_rows(
    app_and_engine: tuple[object, object],
    tmp_path: Path,
) -> None:
    app, engine = app_and_engine
    with Session(engine) as session:
        actor = AuthService(session).initialize_admin("admin", "correct horse").user
        fund = Fund(standard_name="测试慧达一号私募证券投资基金", product_code="FS0010")
        session.add(fund)
        session.flush()
        service = ImportService.from_settings(session, app.state.settings)
        batch = service.create_batch(SourceType.UPLOAD, actor.id)
        service.receive_upload(
            batch.id, "测试慧达一号净值表.xlsx", _ta_confirmed(), actor.id
        )
        service.complete_batch(batch.id, actor.id)
        session.commit()

        result = process_import_batch(session, batch.id, app.state.settings)
        session.commit()

        assert len(result.created_versions) == 1
        version = session.get(ValuationVersion, result.created_versions[0])
        assert version is not None
        snapshot = session.scalar(
            select(FundDailySnapshot).where(
                FundDailySnapshot.valuation_version_id == version.id
            )
        )
        assert snapshot is not None
        assert abs(snapshot.net_asset_value - Decimal("12300068.31")) < Decimal("0.01")
        share_codes = set(
            session.scalars(
                select(ShareClass.share_code).where(ShareClass.fund_id == fund.id)
            )
        )
        assert share_codes == {"ZB1568", "ZB1569"}
        class_snapshots = list(
            session.scalars(
                select(ShareClassDailySnapshot).where(
                    ShareClassDailySnapshot.valuation_version_id == version.id
                )
            )
        )
        assert len(class_snapshots) == 2
        assert sum(
            (item.net_assets for item in class_snapshots), Decimal(0)
        ) == Decimal("12300068.31")


def test_processor_nav_import_is_idempotent(
    app_and_engine: tuple[object, object],
    tmp_path: Path,
) -> None:
    app, engine = app_and_engine
    with Session(engine) as session:
        actor = AuthService(session).initialize_admin("admin", "correct horse").user
        fund = Fund(standard_name="测试慧达二号私募证券投资基金", product_code="XY8410")
        session.add(fund)
        session.flush()
        service = ImportService.from_settings(session, app.state.settings)
        batch = service.create_batch(SourceType.UPLOAD, actor.id)
        service.receive_upload(
            batch.id, "测试慧达二号净值表.xlsx", _ta_export(), actor.id
        )
        service.complete_batch(batch.id, actor.id)
        session.commit()
        result = process_import_batch(session, batch.id, app.state.settings)
        session.commit()
        assert result.processed_files == 1
        assert len(result.created_versions) == 2

        # Re-processing the same batch must not create a second version for
        # any date; the per-source-file duplicate guard covers new batches.
        result_again = process_import_batch(session, batch.id, app.state.settings)
        session.commit()
        assert result_again.created_versions == ()

        versions = list(
            session.scalars(
                select(ValuationVersion).where(ValuationVersion.fund_id == fund.id)
            )
        )
        assert len(versions) == 2


def test_processor_imports_key_value_announcement(
    app_and_engine: tuple[object, object],
    tmp_path: Path,
) -> None:
    app, engine = app_and_engine
    with Session(engine) as session:
        actor = AuthService(session).initialize_admin("admin", "correct horse").user
        fund = Fund(standard_name="测试星辰私募证券投资基金", product_code="FS0009")
        session.add(fund)
        session.flush()
        service = ImportService.from_settings(session, app.state.settings)
        batch = service.create_batch(SourceType.UPLOAD, actor.id)
        service.receive_upload(
            batch.id, "测试星辰每日净值信息.xlsx", _key_value_announcement(), actor.id
        )
        service.complete_batch(batch.id, actor.id)
        session.commit()

        result = process_import_batch(session, batch.id, app.state.settings)
        session.commit()

        assert result.processed_files == 1
        assert result.published_files == 1
        assert len(result.created_versions) == 1
        version = session.get(ValuationVersion, result.created_versions[0])
        assert version is not None
        snapshot = session.scalar(
            select(FundDailySnapshot).where(
                FundDailySnapshot.valuation_version_id == version.id
            )
        )
        assert snapshot is not None
        # SQLite round-trips numeric columns through a float, so compare the
        # money fields within a cent instead of exactly.
        assert abs(snapshot.unit_nav - Decimal("1.0018")) < Decimal("0.0001")
        assert abs(snapshot.net_asset_value - Decimal("8464943.23")) < Decimal("0.01")
        assert abs(snapshot.total_assets - Decimal("8489619.02")) < Decimal("0.01")
        # The announcement carries no liabilities; they are derived so the
        # balance check passes.
        assert abs(snapshot.total_liabilities - Decimal("24675.79")) < Decimal("0.01")
