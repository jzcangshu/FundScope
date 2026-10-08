from __future__ import annotations

import zipfile
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from app.parser.excel_reader import WorkbookLimitError, read_workbook


def _inline_string_sheet(rows: list[list[object]]) -> str:
    """A sheet whose cells are inline strings with a bogus A1 dimension."""

    parts: list[str] = []
    for row_number, row in enumerate(rows, start=1):
        cells = "".join(
            f'<c r="{get_column_letter(c)}{row_number}" t="inlineStr"><is><t>{v}</t></is></c>'
            for c, v in enumerate(row, start=1)
        )
        parts.append(f'<row r="{row_number}">{cells}</row>')
    body = "".join(parts)
    return (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/'
        'spreadsheetml/2006/main">'
        '<dimension ref="A1"/><sheetData>'
        f"{body}</sheetData></worksheet>"
    )


def _save_inline_string_workbook(path: Path, rows: list[list[object]]) -> None:
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    workbook.save(path)
    buffer = BytesIO()
    with (
        zipfile.ZipFile(path) as source,
        zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                content = _inline_string_sheet(rows).encode("utf-8")
            target.writestr(item, content)
    path.write_bytes(buffer.getvalue())


def _save_workbook(
    path: Path, *, sheets: int = 1, rows: int = 1, columns: int = 1
) -> None:
    workbook = Workbook()
    workbook.active.title = "Sheet 1"
    for sheet_index in range(1, sheets):
        workbook.create_sheet(f"Sheet {sheet_index + 1}")
    for sheet in workbook.worksheets:
        for row in range(1, rows + 1):
            for column in range(1, columns + 1):
                sheet.cell(row=row, column=column, value="value")
    workbook.save(path)


def test_xlsx_inline_strings_are_read_despite_bogus_dimension(
    tmp_path: Path,
) -> None:
    # Provider exports declare dimension="A1" and store cells as inline
    # strings; read-only parsing truncates both to the first row.
    path = tmp_path / "inline-strings.xlsx"
    _save_inline_string_workbook(
        path,
        [
            ["净值日期", "产品代码", "产品名称", "单位净值"],
            ["2026-09-15", "FS0005", "测试紫微私募证券投资基金", "2.059"],
            ["2026-09-16", "FS0005", "测试紫微私募证券投资基金", "2.102"],
        ],
    )

    worksheets = read_workbook(path)

    assert len(worksheets) == 1
    rows = worksheets[0].rows
    assert len(rows) == 3
    assert rows[1][1] == "FS0005"
    assert rows[2][3] == "2.102"


def test_xlsx_sheet_count_limit_is_enforced(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "too-many-sheets.xlsx"
    _save_workbook(path, sheets=2)
    monkeypatch.setattr("app.parser.excel_reader.MAX_WORKSHEETS", 1)

    with pytest.raises(WorkbookLimitError, match="workbook_too_many_worksheets"):
        read_workbook(path)


@pytest.mark.parametrize(
    ("rows", "columns", "limit_name", "error_code"),
    [
        (3, 1, "MAX_ROWS_PER_WORKSHEET", "workbook_too_many_rows"),
        (1, 3, "MAX_COLUMNS_PER_WORKSHEET", "workbook_too_many_columns"),
    ],
)
def test_xlsx_row_and_column_limits_are_enforced(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    rows: int,
    columns: int,
    limit_name: str,
    error_code: str,
) -> None:
    path = tmp_path / "oversized.xlsx"
    _save_workbook(path, rows=rows, columns=columns)
    monkeypatch.setattr(f"app.parser.excel_reader.{limit_name}", 2)

    with pytest.raises(WorkbookLimitError, match=error_code):
        read_workbook(path)
